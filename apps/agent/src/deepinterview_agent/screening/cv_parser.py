"""CV parser: PDF/DOCX -> structured text + candidate data.

Extracts raw text from CV documents and uses GPT-4o-mini (cheap, fast) to
produce structured fields (education, experience, skills) from the raw text.
"""

from __future__ import annotations

import re
from typing import Any

from ..core.logging import get_logger

log = get_logger(__name__)


def clean_text(text: str) -> str:
    """Remove excess whitespace and normalize line breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def extract_pdf_text(data: bytes) -> str:
    """Extract text from PDF bytes using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        log.warning("pypdf not installed: %s", exc)
        return ""

    import io
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:
            log.warning("PDF page extraction failed: %s", exc)
    return clean_text("\n".join(parts))


def extract_docx_text(data: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        from docx import Document
    except ImportError as exc:
        log.warning("python-docx not installed: %s", exc)
        return ""

    import io
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return clean_text("\n".join(parts))


def extract_text(data: bytes, mime_type: str) -> str:
    """Dispatch to the right extractor based on MIME type."""
    mime = (mime_type or "").split(";")[0].strip().lower()
    if mime == "application/pdf":
        return extract_pdf_text(data)
    if mime in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return extract_docx_text(data)
    if mime == "text/plain":
        return clean_text(data.decode("utf-8", errors="replace"))
    log.warning("Unknown MIME type for CV: %s", mime)
    return ""


async def parse_structured(cv_text: str) -> dict[str, Any]:
    """Use GPT-4o-mini to extract structured fields from CV text."""
    if not cv_text or len(cv_text.strip()) < 50:
        return {"education": [], "experience": [], "skills": []}

    prompt = f"""Extract structured information from this resume. Respond ONLY with valid JSON:

{{
  "education": [{{"degree": "...", "institution": "...", "year": "..."}}],
  "experience": [{{"title": "...", "company": "...", "duration": "..."}}],
  "skills": ["skill1", "skill2"]
}}

Resume text:
{cv_text[:8000]}"""

    try:
        from ..core.adapters.llm import get_llm
        from ..core.config import get_settings
        adapter = get_llm(get_settings())
        response = await adapter.complete_text(
            system="You are a resume parser. Respond ONLY with valid JSON.",
            user=prompt,
        )
        import json
        return json.loads(response)
    except Exception as exc:
        log.warning("Structured CV parsing failed: %s", exc)
        return {"education": [], "experience": [], "skills": []}


async def parse_cv(data: bytes, mime_type: str) -> dict[str, Any]:
    """Parse CV bytes into structured data.

    Returns: {{ raw_text: str, structured: {{ education[], experience[], skills[] }} }}
    """
    raw_text = extract_text(data, mime_type)
    structured = await parse_structured(raw_text)
    return {
        "raw_text": raw_text,
        "structured": structured,
    }